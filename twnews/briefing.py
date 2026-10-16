"""Word (.docx) briefing generation with python-docx, plus retention pruning."""

from __future__ import annotations

import logging
import re
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .textutil import clean_xml_text, is_chinese, to_simplified

log = logging.getLogger("twnews.briefing")

KIND_LABEL = {"regular": "定时简报", "flash": "快讯简报", "sample": "测试简报"}
FILENAME_RE = re.compile(r"^taiwan-news_\d{8}-\d{6}_(regular|flash|sample)(-\d+)?\.docx$")


@dataclass
class BriefItem:
    source_name: str
    title: str          # verbatim headline
    link: str
    published: datetime | None
    fetched: datetime
    summary: str        # raw summary (any script); converted at render time
    topic: str = ""
    keywords: list[str] = field(default_factory=list)


def _set_fonts(document: docx.Document, font: str) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.name = font
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), font)
        # Strip theme fonts so the explicit font wins.
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
    document.styles["Normal"].font.size = Pt(10.5)


def add_hyperlink(paragraph, url: str, text: str, *, bold: bool = True, size: Pt | None = None):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.underline = True
    run.font.color.rgb = RGBColor(0x0B, 0x57, 0xD0)
    if size:
        run.font.size = size
    hyperlink.append(run._r)  # moves the run inside <w:hyperlink>
    paragraph._p.append(hyperlink)
    return hyperlink


def _fmt(dt: datetime | None) -> str:
    return dt.strftime("%Y-%m-%d %H:%M") if dt else "时间未知"


def _summary_zh(item: BriefItem) -> str:
    if not item.summary:
        return "（该来源未提供摘要）"
    if is_chinese(item.summary):
        return to_simplified(item.summary)
    return f"（英文原文摘要）{item.summary}"


def _group(items: list[BriefItem], group_by: str, source_order: list[str], topic_order: list[str]):
    groups: "OrderedDict[str, list[BriefItem]]" = OrderedDict()
    order = source_order if group_by == "source" else topic_order + ["其他"]
    for key in order:
        groups[key] = []
    for it in items:
        key = it.source_name if group_by == "source" else (it.topic or "其他")
        groups.setdefault(key, []).append(it)
    for key in groups:
        groups[key].sort(key=lambda i: i.published or i.fetched, reverse=True)
    return OrderedDict((k, v) for k, v in groups.items() if v)


def build_briefing(items: list[BriefItem], *, kind: str, generated_at: datetime, out_dir: str,
                   title: str, group_by: str, font: str, source_order: list[str],
                   topic_order: list[str], triggers: list[str] | None = None,
                   window: tuple[datetime, datetime] | None = None) -> Path:
    document = docx.Document()
    _set_fonts(document, font)
    section = document.sections[0]
    document.core_properties.title = f"{title} {generated_at:%Y-%m-%d %H:%M}"
    document.core_properties.author = "TwNewsMonitor"

    label = KIND_LABEL.get(kind, kind)
    heading = document.add_heading(level=0)
    heading.add_run(f"{title}·{label}")
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = sub.add_run(f"{generated_at:%Y年%m月%d日 %H:%M}（台北时间）")
    r.bold = True
    r.font.size = Pt(12)

    groups = _group(items, group_by, source_order, topic_order)
    meta = document.add_paragraph()
    meta_text = f"共收录 {len(items)} 条新闻，涉及 {len(groups)} 个{'来源' if group_by == 'source' else '主题'}。"
    if window:
        meta_text += f" 时间范围：{_fmt(window[0])} 至 {_fmt(window[1])}。"
    meta_text += " 标题保留原文，摘要已转为简体中文。"
    meta.add_run(meta_text).font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    if triggers:
        p = document.add_paragraph()
        p.add_run("触发关键词：").bold = True
        tr = p.add_run("、".join(triggers))
        tr.bold = True
        tr.font.color.rgb = RGBColor(0xC0, 0x1B, 0x1B)

    if not items:
        document.add_paragraph("本时段内没有新的新闻。")

    for group_name, group_items in groups.items():
        document.add_heading(f"{group_name}（{len(group_items)}）", level=1)
        for it in group_items:
            hp = document.add_paragraph(style="List Number")
            hp.paragraph_format.space_after = Pt(2)
            add_hyperlink(hp, it.link, clean_xml_text(it.title), size=Pt(11.5))

            sp = document.add_paragraph()
            sp.paragraph_format.left_indent = Pt(18)
            sp.paragraph_format.space_after = Pt(1)
            sp.add_run(clean_xml_text(_summary_zh(it)))

            mp = document.add_paragraph()
            mp.paragraph_format.left_indent = Pt(18)
            mp.paragraph_format.space_after = Pt(8)
            meta_bits = [f"发布：{_fmt(it.published)}"]
            if it.published is None:
                meta_bits.append(f"抓取：{_fmt(it.fetched)}")
            meta_bits.append(f"来源：{it.source_name}" if group_by == "topic" else f"主题：{it.topic or '其他'}")
            if it.keywords:
                meta_bits.append("关键词：" + "、".join(it.keywords))
            mr = mp.add_run(" · ".join(meta_bits))
            mr.font.size = Pt(9)
            mr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    footer = section.footer.paragraphs[0]
    footer.text = f"TwNewsMonitor 自动生成 · {generated_at:%Y-%m-%d %H:%M:%S} Asia/Taipei"

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"taiwan-news_{generated_at:%Y%m%d-%H%M%S}_{kind}.docx"
    n = 1
    while path.exists():  # two briefings within the same second
        path = out / f"taiwan-news_{generated_at:%Y%m%d-%H%M%S}_{kind}-{n}.docx"
        n += 1
    document.save(path)
    return path


def prune_briefings(out_dir: str, keep_latest: int) -> list[Path]:
    """Keep the newest `keep_latest` files of each kind (regular/flash/sample); delete the rest."""
    removed: list[Path] = []
    if keep_latest <= 0:
        return removed
    by_kind: dict[str, list[Path]] = {}
    for p in Path(out_dir).glob("taiwan-news_*.docx"):
        m = FILENAME_RE.match(p.name)
        if m:
            by_kind.setdefault(m.group(1), []).append(p)
    for files in by_kind.values():
        files.sort(key=lambda p: p.name, reverse=True)
        for old in files[keep_latest:]:
            try:
                old.unlink()
                removed.append(old)
            except OSError as exc:
                log.warning("could not delete old briefing", extra={"file": str(old), "error": str(exc)})
    return removed
