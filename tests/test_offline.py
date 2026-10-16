"""Offline end-to-end tests using httpx.MockTransport (no network needed).

    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import asyncio
import json
import shutil
import sys
import tempfile
import unittest
from datetime import timedelta
from email.utils import format_datetime
from pathlib import Path

import docx
import httpx
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from twnews.briefing import prune_briefings  # noqa: E402
from twnews.config import load_config  # noqa: E402
from twnews.service import Monitor  # noqa: E402
from twnews.textutil import match_keywords, now_taipei, summarize  # noqa: E402


def rss(items: list[tuple[str, str, str]]) -> bytes:
    now = now_taipei()
    body = "".join(
        f"<item><title>{t}</title><link>{link}</link><guid>{link}</guid>"
        f"<description><![CDATA[<p>{d}</p>]]></description>"
        f"<pubDate>{format_datetime(now - timedelta(minutes=i))}</pubDate></item>"
        for i, (t, link, d) in enumerate(items)
    )
    return f'<?xml version="1.0" encoding="UTF-8"?><rss version="2.0"><channel><title>x</title>{body}</channel></rss>'.encode()


class FakeWeb:
    def __init__(self):
        self.feeds: dict[str, bytes] = {}
        self.pages: dict[str, str] = {}
        self.robots: dict[str, str] = {}
        self.requests: list[httpx.Request] = []
        self.push_requests: list[httpx.Request] = []

    def handler(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        url = str(request.url)
        host = request.url.host
        if request.url.path == "/robots.txt":
            if host in self.robots:
                return httpx.Response(200, text=self.robots[host])
            return httpx.Response(404)
        if url in self.feeds:
            etag = f'"{hash(self.feeds[url])}"'
            if request.headers.get("If-None-Match") == etag:
                return httpx.Response(304)
            return httpx.Response(200, content=self.feeds[url], headers={"ETag": etag,
                                                                       "Content-Type": "application/rss+xml"})
        if url in self.pages:
            return httpx.Response(200, text=self.pages[url], headers={"Content-Type": "text/html; charset=utf-8"})
        return httpx.Response(404)

    def push_handler(self, request: httpx.Request) -> httpx.Response:
        self.push_requests.append(request)
        if "day.app" in request.url.host:
            return httpx.Response(200, json={"code": 200, "message": "success"})
        if "telegram" in request.url.host:
            return httpx.Response(200, json={"ok": True})
        return httpx.Response(200, json={"ok": True})


CONFIG = {
    "fetch": {"host_min_gap_seconds": 0, "retries": 1, "backoff_base_seconds": 0.01},
    "sources": [
        {"id": "cna", "name": "中央社", "feeds": ["https://feeds.example.com/cna"]},
        {"id": "ltn", "name": "自由时报", "feeds": ["https://ltn.example.com/broken.xml"],
         "html": {"url": "https://ltn.example.com/list", "link_pattern": r"^https://ltn\.example\.com/news/\d+$"}},
        {"id": "blocked", "name": "Blocked", "feeds": ["https://blocked.example.com/feed"]},
    ],
    "watch_keywords": ["军演", "TSMC"],
    "topics": {"两岸与国防": ["军演", "两岸"], "财经与科技": ["TSMC", "半导体"]},
    "flash": {"cooldown_seconds": 0},
    "links": {"public_base_url": "http://phone.example:8765",
              "file_server": {"enabled": True, "token": "abcdefghijklmnopqrstuvwxyz"}},
    "push": {"channels": [
        {"type": "bark", "device_key": "KEY"},
        {"type": "telegram", "bot_token": "T", "chat_id": "1"},
        {"type": "webhook", "url": "https://hook.example.com/x", "max_file_bytes": 10},
    ]},
}


class OfflineTest(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        (self.dir / "config.yaml").write_text(yaml.safe_dump(CONFIG, allow_unicode=True), encoding="utf-8")
        self.cfg = load_config(self.dir / "config.yaml")
        self.web = FakeWeb()
        self.web.feeds["https://feeds.example.com/cna"] = rss([
            ("國防部：共軍今日在臺海周邊軍演", "https://cna.example.com/news/1?utm_source=rss",
             "國防部表示，共軍今天在臺灣周邊海空域進行聯合演習。國軍嚴密監控。第三句不應出現。"),
            ("台股收盤上漲", "https://cna.example.com/news/2", "台股今天收盤上漲百點。"),
        ])
        self.web.feeds["https://ltn.example.com/broken.xml"] = b"<html>not a feed"
        self.web.pages["https://ltn.example.com/list"] = (
            '<html><body><a href="/news/100">立法院今日三讀通過預算案</a>'
            '<a href="/news/100"><img src="x.jpg"></a><a href="/about">關於我們</a></body></html>')
        self.web.pages["https://ltn.example.com/news/100"] = (
            "<html><head><meta name='description' content='立法院院會今天三讀通過總預算案。'></head>"
            "<body><article><p>立法院院會今天三讀通過總預算案，朝野協商後順利過關。</p></article></body></html>")
        self.web.robots["blocked.example.com"] = "User-agent: *\nDisallow: /\n"

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def monitor(self) -> Monitor:
        return Monitor(self.cfg, transport=httpx.MockTransport(self.web.handler),
                       push_transport=httpx.MockTransport(self.web.push_handler))

    def test_full_cycle(self):
        async def scenario():
            m = self.monitor()
            reports = {r.source_id: r for r in await m.poll_cycle(force=True)}
            self.assertEqual((reports["cna"].fetched, reports["cna"].new, reports["cna"].via), (2, 2, "rss"))
            self.assertEqual(reports["cna"].flash, 0, "first fetch is backlog, no flash")
            self.assertEqual((reports["ltn"].via, reports["ltn"].new), ("html", 1))
            self.assertFalse(reports["blocked"].ok)
            self.assertIn("robots", reports["blocked"].errors[0])
            self.assertFalse(any("blocked.example.com/feed" in str(r.url) for r in self.web.requests))

            # Second poll: conditional GET -> 304, nothing new
            r2 = {r.source_id: r for r in await m.poll_cycle(force=True)}
            self.assertEqual((r2["cna"].via, r2["cna"].new), ("304", 0))
            self.assertTrue(any(r.headers.get("If-None-Match") for r in self.web.requests))

            # New matching item -> flash briefing + push
            self.web.feeds["https://feeds.example.com/cna"] = rss([
                ("TSMC to build new fab", "https://cna.example.com/news/3", "TSMC said on Friday. More later."),
                ("國防部：共軍今日在臺海周邊軍演", "https://cna.example.com/news/1", "dup"),
            ])
            r3 = {r.source_id: r for r in await m.poll_cycle(force=True)}
            self.assertEqual((r3["cna"].new, r3["cna"].flash), (1, 1))
            flashes = list((self.dir / "briefings").glob("*_flash.docx"))
            self.assertEqual(len(flashes), 1)
            self.assertEqual(m.storage.pending_flash(), [])

            # Duplicate content under a different link is dropped by the content hash
            self.web.feeds["https://feeds.example.com/cna"] = rss([
                ("TSMC to build new fab", "https://cna.example.com/news/3-copy", "TSMC said on Friday. More later."),
            ])
            r4 = {r.source_id: r for r in await m.poll_cycle(force=True)}
            self.assertEqual(r4["cna"].new, 0)

            path = await m.regular_briefing()
            self.assertIsNotNone(path)
            text = "\n".join(p.text for p in docx.Document(path).paragraphs)
            self.assertIn("國防部：共軍今日在臺海周邊軍演", text)        # headline verbatim
            self.assertIn("国防部表示，共军今天在台湾周边海空域进行联合演习。", text)  # summary simplified
            self.assertNotIn("第三句", text)                             # 2-sentence cap
            self.assertIn("立法院院会今天三读通过总预算案", text)              # trafilatura enrichment
            self.assertIn("（英文原文摘要）", text)
            rels = [r.target_ref for r in docx.Document(path).part.rels.values() if "hyperlink" in r.reltype]
            self.assertIn("https://cna.example.com/news/1", rels)       # utm stripped, hyperlinked

            await m.close()
            return path

        path = asyncio.run(scenario())
        bark = [r for r in self.web.push_requests if "day.app" in r.url.host]
        body = json.loads(bark[-1].content)
        self.assertEqual(body["url"], f"http://phone.example:8765/abcdefghijklmnopqrstuvwxyz/{path.name}")
        self.assertIn("台湾新闻简报", body["title"])
        tg = [r for r in self.web.push_requests if "telegram" in r.url.host]
        self.assertTrue(any(r.url.path.endswith("/sendDocument") for r in tg))
        hook = [r for r in self.web.push_requests if "hook.example.com" in r.url.host]
        self.assertEqual(json.loads(hook[-1].content)["file_url"].split("/")[-1], path.name)  # too big -> link

    def test_once_mode(self):
        rc = asyncio.run(self.monitor().run_once(push=True))
        self.assertEqual(rc, 0)
        self.assertEqual(len(list((self.dir / "briefings").glob("*_sample.docx"))), 1)
        self.assertTrue(self.web.push_requests)

    def test_helpers(self):
        self.assertEqual(match_keywords("中共軍演逼近", ["军演", "TSMC"]), ["军演"])
        self.assertEqual(match_keywords("tsmc earnings", ["TSMC"]), ["TSMC"])
        self.assertEqual(summarize("第一句。第二句！第三句。"), "第一句。第二句！")
        out = self.dir / "b"
        out.mkdir()
        for i in range(5):
            (out / f"taiwan-news_2026101{i}-120000_regular.docx").write_bytes(b"x")
        (out / "taiwan-news_20261010-120000_flash.docx").write_bytes(b"x")
        prune_briefings(str(out), 2)
        self.assertEqual(sorted(p.name for p in out.iterdir()), [
            "taiwan-news_20261010-120000_flash.docx",
            "taiwan-news_20261013-120000_regular.docx",
            "taiwan-news_20261014-120000_regular.docx"])


if __name__ == "__main__":
    unittest.main()
